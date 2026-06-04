"""IO helpers, problem loading, and rate limiting.

Most functions in this module are extracted from the upstream
`math_prove/main.py` (and `run_parallel_batch.py` for `RpmLimiter`). The
extraction is mechanical — the code is byte-identical, only the import
target has changed to use the vendored `parser` module.
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import threading
import time
from collections import deque
from pathlib import Path
from typing import Any, Deque, Dict, Iterable, List, Optional, Tuple

from shared.vendor.parser import (
    parse_and_validate,
    fallback_solution,
    solution_to_json,
)


ID_FIELDS: Tuple[str, ...] = (
    "problem_id",
    "id",
    "question_id",
    "qid",
    "uid",
    "index",
)
TEXT_FIELDS: Tuple[str, ...] = (
    "problem_text",
    "problem",
    "question",
    "text",
    "content",
    "题目",
)

SAMPLE_PROBLEMS: List[Dict[str, str]] = [
    {
        "problem_id": "demo_001",
        "problem_text": "Find all real roots of x^4 - 5x^2 + 4 = 0.",
    },
    {
        "problem_id": "demo_002",
        "problem_text": "Given f(z) = (z^2 + 1)/(z - i), find the residue at z = i.",
    },
    {
        "problem_id": "demo_003",
        "problem_text": (
            "Maximize 3x + 4y subject to x + 2y <= 8, 3x + y <= 9, "
            "x >= 0, y >= 0."
        ),
    },
]


def load_problems(path: str) -> List[Dict[str, Any]]:
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix == ".jsonl":
        rows = _load_jsonl(source)
    elif suffix == ".json":
        rows = _load_json(source)
    elif suffix == ".csv":
        rows = _load_csv(source)
    elif suffix in {".xlsx", ".xls"}:
        rows = _load_excel(source)
    else:
        raise ValueError(f"Unsupported input format: {source.suffix}")

    problems = [_normalize_problem_row(row, idx) for idx, row in enumerate(rows, start=1)]
    return problems


def _load_jsonl(path: Path) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as handle:
        for line_no, line in enumerate(handle, start=1):
            line = line.strip()
            if not line:
                continue
            obj = json.loads(line)
            if not isinstance(obj, dict):
                raise ValueError(f"JSONL line {line_no} is not an object")
            rows.append(obj)
    return rows


def _load_json(path: Path) -> List[Dict[str, Any]]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(data, list):
        rows = data
    elif isinstance(data, dict):
        for key in ("problems", "data", "items", "questions"):
            if isinstance(data.get(key), list):
                rows = data[key]
                break
        else:
            rows = [data]
    else:
        raise ValueError("JSON input must be an object or array")
    if not all(isinstance(row, dict) for row in rows):
        raise ValueError("JSON problem rows must be objects")
    return list(rows)


def _load_csv(path: Path) -> List[Dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def _load_excel(path: Path) -> List[Dict[str, Any]]:
    try:
        import pandas as pd  # type: ignore

        return pd.read_excel(path).fillna("").to_dict(orient="records")
    except ImportError:
        pass

    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError as exc:
        raise ImportError("Reading Excel requires pandas or openpyxl") from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(cell or "").strip() for cell in rows[0]]
    data = []
    for row in rows[1:]:
        data.append({headers[i]: row[i] if i < len(row) else "" for i in range(len(headers))})
    return data


def _normalize_problem_row(row: Dict[str, Any], index: int) -> Dict[str, Any]:
    pid = _first_present(row, ID_FIELDS)
    text = _first_present(row, TEXT_FIELDS)
    if pid is None:
        pid = f"{index:03d}"
    if text is None:
        text = ""

    pid_str = str(pid).strip() or f"{index:03d}"
    text_str = str(text).strip()
    metadata = {key: value for key, value in row.items() if key not in ID_FIELDS + TEXT_FIELDS}
    return {
        "problem_id": pid_str,
        "problem_text": text_str,
        "raw_metadata": metadata,
    }


def _first_present(row: Dict[str, Any], fields: Iterable[str]) -> Optional[Any]:
    lower_map = {str(key).lower(): key for key in row.keys()}
    for field in fields:
        if field in row and row[field] not in (None, ""):
            return row[field]
        key = lower_map.get(field.lower())
        if key is not None and row[key] not in (None, ""):
            return row[key]
    return None


def read_existing_results(path: Path) -> Dict[str, Dict[str, Any]]:
    """Read a JSONL result file into a {problem_id: row} dict for resume support."""
    results: Dict[str, Dict[str, Any]] = {}
    if not path.exists():
        return results
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if not line:
                continue
            try:
                obj = json.loads(line)
                pid = str(obj.get("problem_id", "")).strip()
                if pid:
                    results[pid] = obj
            except Exception:
                continue
    return results


def load_and_validate_results(path: Path) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """Re-parse every JSONL line as MathSolution; collect schema errors per line."""
    results: List[Dict[str, Any]] = []
    errors: List[Dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as handle:
        for line_no, line in enumerate(handle, start=1):
            line = line.strip()
            if not line:
                continue
            try:
                obj = json.loads(line)
                pid = str(obj.get("problem_id") or f"line_{line_no}")
                solution = parse_and_validate(json.dumps(obj, ensure_ascii=False), pid)
                results.append(solution.model_dump(mode="json"))
            except Exception as exc:
                errors.append({"line": line_no, "error": f"{type(exc).__name__}: {exc}"})
    return results, errors


def write_problem_log(log_dir: Path, problem_id: str, run_log: Dict[str, Any]) -> None:
    """Write per-problem log JSON; sanitizes the problem_id for the filename."""
    safe_id = re.sub(r"[^A-Za-z0-9_.-]+", "_", str(problem_id)) or "unknown"
    path = log_dir / f"{safe_id}.json"
    path.write_text(json.dumps(run_log, ensure_ascii=False, indent=2), encoding="utf-8")


DEFAULT_API_BASE = "https://api.openai.com/v1/chat/completions"


def resolve_api_config(args: argparse.Namespace) -> Tuple[Optional[str], Optional[str]]:
    """Resolve API key + base URL from argparse Namespace with env fallbacks.

    Reads:
      - args.api_key (or "OPENAI_API_KEY" env var)
      - args.api_base (or "LLM_API_BASE" env var, default: OpenAI chat-completions)
    """
    api_key = getattr(args, "api_key", None) or os.environ.get("OPENAI_API_KEY")
    api_base = (
        getattr(args, "api_base", None)
        or os.environ.get("LLM_API_BASE")
        or DEFAULT_API_BASE
    )
    return api_key, api_base


class RpmLimiter:
    """Thread-safe sliding-window request limiter (60-second window).

    The limiter is acquired around each LLM call so the rate is enforced on
    API requests rather than on problem-level tasks. Safe to share across
    multiple worker threads.
    """

    def __init__(self, rpm_limit: int) -> None:
        self.rpm_limit = max(1, int(rpm_limit))
        self._events: Deque[float] = deque()
        self._lock = threading.Lock()

    def acquire(self) -> None:
        while True:
            with self._lock:
                now = time.monotonic()
                while self._events and now - self._events[0] >= 60.0:
                    self._events.popleft()
                if len(self._events) < self.rpm_limit:
                    self._events.append(now)
                    return
                wait_seconds = max(0.05, 60.0 - (now - self._events[0]))
            time.sleep(wait_seconds)


__all__ = [
    "ID_FIELDS",
    "TEXT_FIELDS",
    "SAMPLE_PROBLEMS",
    "load_problems",
    "read_existing_results",
    "load_and_validate_results",
    "write_problem_log",
    "resolve_api_config",
    "RpmLimiter",
    "fallback_solution",
    "solution_to_json",
]
